"""openpyxl を用いた Excel (.xlsx) 生成モジュール。

ハイライトテキストに数値が複数含まれる場合（＝表をハイライトした場合）、
数値トークンを 1 セルずつ別カラム（値1, 値2, …）へ展開する。
通貨記号・カンマ・括弧（会計上の負数）・%・倍率(x) を解釈し、
純粋な数値は Excel の数値型として書き込む。
"""

from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 先頭固定列
BASE_HEADERS = ["ハイライトテキスト", "メモ", "PDFファイル名", "抽出日時"]

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_STRIPE_FILL = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
_WRAP = Alignment(wrap_text=True, vertical="top")
_RIGHT = Alignment(horizontal="right", vertical="top")

_CURRENCY = "$€£¥"
_DASHES = {"-", "–", "—"}

# この数以上の数値トークンを含むハイライトのみ「表」とみなして列分割する。
# （散文中にたまたま数字が混ざる場合に列がばらけるのを防ぐ）
_SPLIT_THRESHOLD = 3


def _safe_sheet_title(filename: str) -> str:
    """ファイル名からシート名を生成する（拡張子除去・禁止文字置換・31文字制限）。"""
    base = re.sub(r"\.pdf$", "", filename, flags=re.IGNORECASE)
    base = f"{base}_Highlights"
    base = re.sub(r"[:\\/?*\[\]]", "_", base)  # Excel のシート名禁止文字
    return base[:31] or "Highlights"


def _tokenize(text: str) -> list[str]:
    """空白で分割し、単独の通貨記号を直後のトークンへ結合する（"$" "121,422" → "$121,422"）。"""
    raw = text.split()
    tokens: list[str] = []
    i = 0
    while i < len(raw):
        t = raw[i]
        if t in _CURRENCY and i + 1 < len(raw):
            tokens.append(t + raw[i + 1])
            i += 2
        else:
            tokens.append(t)
            i += 1
    return tokens


def _is_numeric_token(tok: str) -> bool:
    """通貨記号・カンマ・括弧・%・倍率(x) を除いて数値とみなせるか。"""
    s = tok.strip(_CURRENCY + "()%")
    s = s.replace(",", "")
    if s[-1:] in ("x", "X"):  # 2.5x のような倍率表記
        s = s[:-1]
    if not s or s in _DASHES:
        return False
    return bool(re.fullmatch(r"-?\d*\.?\d+", s))


def _parse_cell(tok: str):
    """トークンを Excel セル値へ変換する。

    純粋な数値（通貨・カンマ・会計負数 (123) を含む）は int/float に変換。
    %・倍率(x)・n/a・ダッシュなどは元の文字列のまま返す。
    """
    neg = tok.startswith("(") and tok.endswith(")")
    s = tok[1:-1] if neg else tok
    s = s.lstrip(_CURRENCY).strip().replace(",", "")
    if re.fullmatch(r"\d*\.?\d+", s):
        val = float(s)
        val = -val if neg else val
        return int(val) if val.is_integer() else val
    return tok


def split_highlight(text: str):
    """ハイライト本文を (ラベル, 値リスト) に分割する。

    先頭の連続した非数値トークンをラベル（1セル）にまとめ、
    以降のトークンを 1 つずつ値セルにする。数値が少ない（散文）の場合は
    全文をラベルにして分割しない。
    """
    text = text or ""
    tokens = _tokenize(text)
    numeric_count = sum(_is_numeric_token(t) for t in tokens)
    if numeric_count < _SPLIT_THRESHOLD:
        return text.strip(), []

    i = 0
    label_parts: list[str] = []
    while i < len(tokens) and not _is_numeric_token(tokens[i]):
        label_parts.append(tokens[i])
        i += 1
    label = " ".join(label_parts)
    values = [_parse_cell(t) for t in tokens[i:]]
    return label, values


def build_table(highlights: list[dict], filename: str):
    """Excel / Google Sheets 共通の (ヘッダー, 行リスト) を生成する。

    行は [ラベル, メモ, ファイル名, 抽出日時, 値1, 値2, …]。
    値セルは int / float / str の混在（Sheets でもそのまま使える）。
    """
    extracted_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    parsed = []
    max_values = 0
    for h in highlights:
        cells = h.get("cells")
        if cells:
            # 表として再構成済み：先頭セル=ラベル、残り=値（数値変換を試みる）
            label = cells[0] if cells else ""
            values = [_parse_cell(c) for c in cells[1:]]
        else:
            label, values = split_highlight(h.get("text", ""))
        max_values = max(max_values, len(values))
        parsed.append((label, h.get("note", ""), values))

    headers = BASE_HEADERS + [f"値{i + 1}" for i in range(max_values)]
    rows = [
        [label, note, filename, extracted_at, *values]
        for label, note, values in parsed
    ]
    return headers, rows


def build_excel(highlights: list[dict], filename: str) -> bytes:
    """ハイライト一覧から xlsx バイト列を生成する。"""
    headers, rows = build_table(highlights, filename)
    n_cols = len(headers)

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(filename)

    # ヘッダー行
    ws.append(headers)
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            vertical="center", horizontal="left" if col_idx <= 4 else "right"
        )

    # データ行
    for i, row in enumerate(rows):
        ws.append(row)
        excel_row = i + 2
        stripe = i % 2 == 1
        for col_idx in range(1, n_cols + 1):
            cell = ws.cell(row=excel_row, column=col_idx)
            if stripe:
                cell.fill = _STRIPE_FILL
            if col_idx in (1, 2):  # ラベル・メモは折り返し
                cell.alignment = _WRAP
            elif col_idx >= 5:  # 値列
                cell.alignment = _RIGHT
                if isinstance(cell.value, bool):
                    pass
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"

    # 列幅
    for col_idx in range(1, n_cols + 1):
        letter = get_column_letter(col_idx)
        if col_idx == 1:
            content = [len(headers[0])] + [
                max((len(line) for line in str(r[0]).splitlines()), default=0)
                for r in rows
            ]
            ws.column_dimensions[letter].width = min(max(content) + 2, 60)
        elif col_idx == 2:
            ws.column_dimensions[letter].width = 30
        elif col_idx == 3:
            ws.column_dimensions[letter].width = 28
        elif col_idx == 4:
            ws.column_dimensions[letter].width = 18
        else:
            ws.column_dimensions[letter].width = 14

    ws.freeze_panes = "A2"  # ヘッダー行を固定

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
